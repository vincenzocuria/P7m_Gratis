from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from io import BytesIO
from typing import List, Dict, Any


class ExcelGenerator:
    """Generatore di file Excel per riepilogo fatture."""
    
    def generate_summary(self, fatture: List[Dict[str, Any]]) -> bytes:
        """
        Genera un file Excel con il riepilogo delle fatture.
        
        Args:
            fatture: Lista di dizionari con i dati delle fatture
        
        Returns:
            Bytes del file Excel generato
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "Riepilogo Fatture"
        
        headers = [
            'Fornitore',
            'P.IVA Fornitore',
            'Numero',
            'Data',
            'Tipo Documento',
            'Imponibile',
            'IVA',
            'Totale',
            'Divisa'
        ]
        
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        row_num = 2
        for fattura in fatture:
            cedente = fattura.get('cedente', {})
            dati_gen = fattura.get('dati_generali', {})
            riepilogo = fattura.get('riepilogo', [])
            
            totale_imponibile = sum(
                float(item.get('imponibile', '0').replace(',', '.'))
                for item in riepilogo
            )
            totale_iva = sum(
                float(item.get('imposta', '0').replace(',', '.'))
                for item in riepilogo
            )
            totale = totale_imponibile + totale_iva
            
            tipo_doc_map = {
                'TD01': 'Fattura',
                'TD02': 'Acconto/Anticipo su fattura',
                'TD03': 'Acconto/Anticipo su parcella',
                'TD04': 'Nota di credito',
                'TD05': 'Nota di debito',
                'TD06': 'Parcella',
            }
            tipo_doc = tipo_doc_map.get(dati_gen.get('tipo_documento', ''), dati_gen.get('tipo_documento', ''))
            
            ws.cell(row=row_num, column=1, value=cedente.get('denominazione', ''))
            ws.cell(row=row_num, column=2, value=cedente.get('partita_iva', ''))
            ws.cell(row=row_num, column=3, value=dati_gen.get('numero', ''))
            ws.cell(row=row_num, column=4, value=self._format_date_excel(dati_gen.get('data', '')))
            ws.cell(row=row_num, column=5, value=tipo_doc)
            ws.cell(row=row_num, column=6, value=totale_imponibile)
            ws.cell(row=row_num, column=7, value=totale_iva)
            ws.cell(row=row_num, column=8, value=totale)
            ws.cell(row=row_num, column=9, value=dati_gen.get('divisa', 'EUR'))
            
            for col in [6, 7, 8]:
                ws.cell(row=row_num, column=col).number_format = '#,##0.00'
            
            row_num += 1
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width
        
        ws2 = wb.create_sheet("Riepilogo IVA")
        self._create_iva_summary(ws2, fatture)
        
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        
        return output.getvalue()
    
    def _create_iva_summary(self, ws, fatture):
        """Crea un foglio di riepilogo IVA per aliquota."""
        headers = ['Aliquota IVA', 'Imponibile Totale', 'IVA Totale']
        
        header_fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        
        aliquote = {}
        
        for fattura in fatture:
            riepilogo = fattura.get('riepilogo', [])
            for item in riepilogo:
                aliquota = item.get('aliquota_iva', '0')
                imponibile = float(item.get('imponibile', '0').replace(',', '.'))
                imposta = float(item.get('imposta', '0').replace(',', '.'))
                
                if aliquota not in aliquote:
                    aliquote[aliquota] = {'imponibile': 0.0, 'imposta': 0.0}
                
                aliquote[aliquota]['imponibile'] += imponibile
                aliquote[aliquota]['imposta'] += imposta
        
        row_num = 2
        for aliquota in sorted(aliquote.keys(), key=lambda x: float(x) if x else 0):
            ws.cell(row=row_num, column=1, value=f"{aliquota}%")
            ws.cell(row=row_num, column=2, value=aliquote[aliquota]['imponibile'])
            ws.cell(row=row_num, column=3, value=aliquote[aliquota]['imposta'])
            
            for col in [2, 3]:
                ws.cell(row=row_num, column=col).number_format = '#,##0.00'
            
            row_num += 1
        
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 30)
            ws.column_dimensions[column].width = adjusted_width
    
    def _format_date_excel(self, date_str: str) -> str:
        """Formatta una data per Excel."""
        if not date_str or len(date_str) != 10:
            return date_str
        try:
            parts = date_str.split('-')
            return f"{parts[2]}/{parts[1]}/{parts[0]}"
        except:
            return date_str
